from __future__ import annotations

import io
import logging
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from core.settings import get_settings
from services.database import get_db
from services.llm_factory import get_llm_gateway
from core.skill_engine.base import SkillContext

logger = logging.getLogger(__name__)

router = APIRouter()

UPLOAD_DIR = Path("./uploads/formatted")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

TEMPLATE_DIR = Path("./templates")
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)

EXPORT_DIR = Path("./uploads/exports")
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

MAX_UPLOAD_BYTES = 50 * 1024 * 1024

RE_MD_HEADING = re.compile(r'^(#{1,6})\s+(.*)')
RE_MD_TABLE_SEP = re.compile(r'^\|?[\s:|-]+$')
RE_MD_BOLD = re.compile(r'\*\*(.+?)\*\*')


def _add_md_runs(para, text: str):
    """把行内 **加粗** 语法拆成加粗/普通 run,其余文本原样写入。"""
    pos = 0
    for m in RE_MD_BOLD.finditer(text):
        if m.start() > pos:
            para.add_run(text[pos:m.start()])
        r = para.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        para.add_run(text[pos:])


_PPR_TAG_ORDER = (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
)


def _attach_heading_auto_numbering(doc):
    """Heading 1-5 样式绑定多级列表自动编号(WPS/Word 标题自动编号):
    章标题 1/2/…、小节 1.1/1.1.1 由文档编号引擎维护,增删章节自动重排,
    不在标题文本里手写编号。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = doc.part.numbering_part.element
    abs_ids = [int(a.get(qn("w:abstractNumId")))
               for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = (max(abs_ids) + 1) if abs_ids else 0
    num_id = (max(num_ids) + 1) if num_ids else 1

    style_ids = ["Heading1", "Heading2", "Heading3", "Heading4", "Heading5"]
    style_names = ["Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"]
    lvl_texts = ["%1", "%1.%2", "%1.%2.%3", "%1.%2.%3.%4", "%1.%2.%3.%4.%5"]

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "multilevel")
    abstract.append(mlt)
    for ilvl in range(5):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), "decimal")
        lvl.append(fmt)
        # w:lvl 子元素顺序: start, numFmt, pStyle, suff, lvlText, lvlJc, pPr
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), style_ids[ilvl])
        lvl.append(pstyle)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")  # 编号与标题间用单个空格,不用制表位
        lvl.append(suff)
        lt = OxmlElement("w:lvlText")
        lt.set(qn("w:val"), lvl_texts[ilvl])
        lvl.append(lt)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:firstLine"), "0")
        ppr.append(ind)
        lvl.append(ppr)
        abstract.append(lvl)

    # schema 要求所有 abstractNum 位于 num 之前: 插到第一个 num 前,没有则追加
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    aref = OxmlElement("w:abstractNumId")
    aref.set(qn("w:val"), str(abstract_id))
    num.append(aref)
    numbering.append(num)

    for ilvl, (sid, sname) in enumerate(zip(style_ids, style_names)):
        style = doc.styles[sname]
        assert style.style_id == sid, f"unexpected style id {style.style_id} for {sname}"
        ppr = style.element.get_or_add_pPr()
        numpr = ppr.find(qn("w:numPr"))
        if numpr is None:
            numpr = OxmlElement("w:numPr")
            my_idx = _PPR_TAG_ORDER.index("numPr")
            inserted = False
            for child in ppr:
                tag = child.tag.rsplit("}", 1)[-1]
                if tag in _PPR_TAG_ORDER and _PPR_TAG_ORDER.index(tag) > my_idx:
                    child.addprevious(numpr)
                    inserted = True
                    break
            if not inserted:
                ppr.append(numpr)
        ilvl_el = OxmlElement("w:ilvl")
        ilvl_el.set(qn("w:val"), str(ilvl))
        numid_el = OxmlElement("w:numId")
        numid_el.set(qn("w:val"), str(num_id))
        numpr.append(ilvl_el)
        numpr.append(numid_el)


def _append_markdown_lines(doc, content: str):
    """章节 markdown 正文 → docx 元素: # 标题 / | 表格 | / **加粗** / 普通段落。

    标题编号由 Heading 样式绑定的多级列表自动生成,这里只剥离 LLM 沿用的
    原编号(如第 4 章内容写 4.1),避免与自动编号重复。
    """
    lines = [ln.strip() for ln in content.split("\n")]
    # 层级归一化: LLM 常直接用 ### 写一级小节(跳过 ##),统计本章实际最外层
    # 小节级别并整体上移,保证最外层小节落在 Heading2,自动编号才能是 1.1 而非 1.1.1
    md_levels = [min(len(m.group(1)), 5) for ln in lines
                 if (m := RE_MD_HEADING.match(ln))]
    base = min((lv for lv in md_levels if lv >= 2), default=2)
    shift = base - 2
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1
            continue

        # markdown 表格: 表头行 + |---| 分隔行 + 数据行 → 真 docx 表格
        if (
            line.startswith("|") and line.endswith("|")
            and i + 1 < len(lines)
            and RE_MD_TABLE_SEP.match(lines[i + 1]) and "-" in lines[i + 1]
        ):
            rows = [line]
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append(lines[j])
                j += 1
            header = [c.strip() for c in rows[0].strip("|").split("|")]
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for k, val in enumerate(header):
                table.rows[0].cells[k].text = val
            for raw in rows[1:]:
                cells = [c.strip() for c in raw.strip("|").split("|")]
                row = table.add_row()
                for k in range(len(header)):
                    row.cells[k].text = cells[k] if k < len(cells) else ""
            doc.add_paragraph()
            i = j
            continue

        # 标题: 先匹配长 # 串(##### 先于 ####),级别归一化后映射到 docx heading
        hm = RE_MD_HEADING.match(line)
        if hm:
            level = min(len(hm.group(1)), 5)
            if level <= 1:
                level = 2  # 章内 # 视为一级小节(章标题由章节循环添加)
            level = min(max(level - shift, 2), 5)
            title_text = hm.group(2).strip()
            if level >= 2:
                title_text = re.sub(r'^\d+(\.\d+)*\s+', '', title_text)
            doc.add_heading(title_text, level=level)
        else:
            _add_md_runs(doc.add_paragraph(), line)
        i += 1


def _save_upload(file: UploadFile, prefix: str = "") -> Path:
    if file.size and file.size > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    name = Path(file.filename or "upload").name
    safe = "".join(c for c in name if c.isalnum() or c in ("-", "_", "."))
    dest = UPLOAD_DIR / f"{prefix}{safe}"
    content = file.file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_UPLOAD_BYTES // 1024 // 1024}MB 上限")
    dest.write_bytes(content)
    return dest


@router.post("/format")
async def format_document(
    file: UploadFile = File(...),
    template: str = "default",
    mode: str = Query("format", description="format|check|diff|beautify"),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file)

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": mode,
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success and skill_result.data:
        result_data = skill_result.data
        output_path = result_data.get("output_path", "")
        if output_path and Path(output_path).exists():
            result_data["file_name"] = Path(output_path).name
            try:
                result_data["file_size"] = Path(output_path).stat().st_size
            except OSError:
                pass

        return {"success": True, "data": result_data}

    return {"success": False, "error": skill_result.error}


@router.post("/check-format")
async def check_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="check_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": "check",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/diff-format")
async def diff_format(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="diff_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "template": template,
            "mode": "diff",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/beautify")
async def beautify_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    temp_path = _save_upload(file, prefix="beautify_")

    from services.format.skills.docx_format_skill import DocxFormatSkill

    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(temp_path),
            "mode": "beautify",
        },
    )
    skill_result = await skill.safe_execute(ctx)

    if skill_result.success and skill_result.data:
        return {"success": True, "data": skill_result.data}
    return {"success": False, "error": skill_result.error}


@router.post("/export-doc")
async def export_to_doc(
    file: UploadFile = File(...),
    template: str = "default",
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 doc 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")

    docx_to_send: Path = src
    if apply_format:
        formatted = await _run_format_skill(src, template, db)
        if formatted:
            docx_to_send = formatted

    from services.format.skills.doc_export_skill import DocExportSkill
    skill = DocExportSkill()
    if not docx_to_send.exists():
        raise HTTPException(status_code=500, detail="排版产物丢失")

    result = skill._try_libreoffice_doc(docx_to_send, docx_to_send.parent)
    if not result.get("success"):
        if docx_to_send != src:
            try:
                docx_to_send.unlink(missing_ok=True)
            except OSError:
                pass
        raise HTTPException(status_code=500, detail=result.get("error", "doc 转换失败"))

    doc_path = Path(result["output_path"])
    data = doc_path.read_bytes()
    try:
        doc_path.unlink(missing_ok=True)
        if docx_to_send != src:
            docx_to_send.unlink(missing_ok=True)
    except OSError:
        pass

    filename = f"{Path(file.filename or 'document').stem}.doc"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/msword",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/export-pdf")
async def export_to_pdf(
    file: UploadFile = File(...),
    template: str = "default",
    apply_format: bool = Form(True),
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版(可选) → 转 pdf 字节流下载。"""
    src = _save_upload(file, prefix="srcdocx_")

    docx_to_send: Path = src
    if apply_format:
        formatted = await _run_format_skill(src, template, db)
        if formatted:
            docx_to_send = formatted

    from services.format.skills.pdf_export_skill import PdfExportSkill
    gateway = get_llm_gateway()
    skill = PdfExportSkill()
    out_dir = docx_to_send.parent
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "input_path": str(docx_to_send),
            "output_dir": str(out_dir),
        },
    )
    skill_result = await skill.safe_execute(ctx)
    if not skill_result.success or not skill_result.data:
        if docx_to_send != src:
            try:
                docx_to_send.unlink(missing_ok=True)
            except OSError:
                pass
        raise HTTPException(status_code=500, detail=skill_result.error or "PDF 转换失败")

    pdf_path = Path(skill_result.data["output_path"])
    data = pdf_path.read_bytes()
    try:
        pdf_path.unlink(missing_ok=True)
        if docx_to_send != src:
            docx_to_send.unlink(missing_ok=True)
    except OSError:
        pass

    filename = f"{Path(file.filename or 'document').stem}.pdf"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/export-formatted-docx")
async def export_formatted_docx(
    file: UploadFile = File(...),
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    """统一入口: 接收 docx → 排版 → 直接下载 docx 字节流。"""
    src = _save_upload(file, prefix="srcdocx_")
    formatted = await _run_format_skill(src, template, db) or src
    data = formatted.read_bytes()
    try:
        if formatted != src:
            formatted.unlink(missing_ok=True)
    except OSError:
        pass
    filename = f"{Path(file.filename or 'document').stem}_formatted.docx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


async def _run_format_skill(src: Path, template: str, db: AsyncSession) -> Path | None:
    from services.format.skills.docx_format_skill import DocxFormatSkill
    gateway = get_llm_gateway()
    skill = DocxFormatSkill()
    ctx = SkillContext(
        project_id="",
        db=db,
        llm=gateway,
        parameters={
            "file_path": str(src),
            "template": template,
            "mode": "format",
        },
    )
    result = await skill.safe_execute(ctx)
    if not result.success or not result.data:
        logger.warning("format skill 失败: %s", result.error)
        return None
    out = result.data.get("output_path", "")
    if out and Path(out).exists():
        return Path(out)
    return None


@router.post("/from-project/{project_id}")
async def format_from_project(
    project_id: str,
    template: str = "default",
    db: AsyncSession = Depends(get_db),
):
    """从项目章节组装为 docx 文件（无需上传）"""
    from sqlalchemy import select
    from services.models import Project, Chapter

    proj_result = await db.execute(select(Project).where(Project.id == project_id))
    project = proj_result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    ch_result = await db.execute(
        select(Chapter).where(Chapter.project_id == project_id).order_by(Chapter.sort_order)
    )
    chapters = ch_result.scalars().all()
    if not chapters:
        raise HTTPException(
            status_code=400,
            detail="项目尚无任何章节大纲，请先到「投标生成」完成大纲生成。",
        )

    empty_chapters = [ch for ch in chapters if not (ch.content and ch.content.strip())]
    if empty_chapters:
        empty_titles = [ch.title for ch in empty_chapters[:5]]
        detail_msg = (
            f"项目存在 {len(empty_chapters)}/{len(chapters)} 个章节尚未生成正文，"
            f"无法执行文档输出。请先到「投标生成」完成正文生成。"
            f"未生成章节示例: {', '.join(empty_titles)}"
        )
        if len(empty_chapters) == len(chapters):
            raise HTTPException(status_code=400, detail=detail_msg)
        if len(empty_chapters) / len(chapters) > 0.5:
            raise HTTPException(
                status_code=400,
                detail=detail_msg + "（空内容章节超过50%，已禁止操作）",
            )

    title = f"{project.name} - 投标文件"

    import tempfile
    from docx import Document
    from docx.shared import Pt
    from pathlib import Path
    from core.skill_engine.base import SkillContext
    from services.format.skills.docx_format_skill import DocxFormatSkill

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    _attach_heading_auto_numbering(doc)
    doc.add_heading(title, level=0)

    for ch in chapters:
        if not (ch.content and ch.content.strip()):
            continue
        h = doc.add_heading(ch.title, level=1)
        # 每个大章从新页开始(标书惯例)
        h.paragraph_format.page_break_before = True
        _append_markdown_lines(doc, ch.content)

    output_dir = Path(tempfile.gettempdir()) / "bidmaster_format"
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[^\w\u4e00-\u9fff\-_]', '_', project.name)[:60] or "bid"
    src_path = output_dir / f"{safe_name}_src_{project_id[:8]}.docx"
    doc.save(str(src_path))

    from services.llm_factory import get_llm_gateway
    gateway = get_llm_gateway()
    skill = DocxFormatSkill()

    ctx = SkillContext(
        project_id=project_id,
        db=db,
        llm=gateway,
        parameters={"file_path": str(src_path), "template": template, "mode": "format"},
    )
    result = await skill.safe_execute(ctx)
    if not result.success:
        try:
            src_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=result.error or "项目章节组装失败")

    final_output = result.data.get("output_path") if isinstance(result.data, dict) else None
    if final_output:
        # skill 产物默认写在临时目录,/api/format/download 仅放行 uploads/formatted,
        # 持久化过去否则前端无法下载
        import asyncio
        import shutil
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        persist_path = UPLOAD_DIR / Path(final_output).name
        await asyncio.to_thread(shutil.copy2, final_output, persist_path)
        final_output = str(persist_path)

    return {
        "success": True,
        "output_path": final_output,
        "project_name": project.name,
        "chapter_count": len(chapters),
        "generated_chapter_count": len(chapters) - len(empty_chapters),
        "empty_chapter_count": len(empty_chapters),
        "chapters": [
            {
                "id": str(c.id),
                "title": c.title,
                "word_count": c.word_count or 0,
                "status": c.status,
                "has_content": bool(c.content and c.content.strip()),
            }
            for c in chapters
        ],
    }


@router.get("/templates")
async def list_templates():
    templates = [
        {"name": "default", "description": "默认标书排版模板(仿宋正文+黑体标题)"},
        {"name": "government", "description": "政府采购标书模板(方正小标宋+仿宋)"},
        {"name": "engineering", "description": "工程标书模板(宋体正文+黑体标题)"},
    ]

    for tpl_file in TEMPLATE_DIR.glob("*.yaml"):
        name = tpl_file.stem
        if not any(t["name"] == name for t in templates):
            templates.append({"name": name, "description": f"自定义模板: {name}"})

    return {"templates": templates}


@router.get("/templates/{template_name}")
async def get_template_detail(template_name: str):
    from services.format.skills.docx_format_skill import DEFAULT_FORMAT_CONFIG

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    if template_path.exists():
        try:
            import yaml
            with open(template_path, encoding="utf-8") as f:
                custom_config = yaml.safe_load(f) or {}
            merged = {**DEFAULT_FORMAT_CONFIG, **custom_config}
        except Exception:
            merged = DEFAULT_FORMAT_CONFIG
    else:
        merged = DEFAULT_FORMAT_CONFIG

    return {"name": template_name, "config": merged}


@router.put("/templates/{template_name}")
async def save_template(template_name: str, config: dict):
    import yaml

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    with open(template_path, "w", encoding="utf-8") as f:
        yaml.dump(config, f, allow_unicode=True, default_flow_style=False)

    return {"success": True, "name": template_name}


@router.delete("/templates/{template_name}")
async def delete_template(template_name: str):
    if template_name in ("default", "government", "engineering"):
        return {"success": False, "error": "内置模板不可删除"}

    template_path = TEMPLATE_DIR / f"{template_name}.yaml"
    if template_path.exists():
        template_path.unlink()
        return {"success": True}
    return {"success": False, "error": "模板不存在"}


@router.get("/download")
async def download_formatted_file(path: str):
    """下载由 /format 产生的产物（限定 uploads/formatted 目录）。"""
    p = Path(path)
    if p.is_absolute():
        p = p.resolve()
    else:
        # skill 返回的 output_path 可能是 "uploads/formatted/xxx.docx" 相对路径,
        # 也可能是纯文件名: 先按工作目录解析,不存在再退回 uploads/formatted 按文件名找
        cand = p.resolve()
        p = cand if cand.is_file() else (UPLOAD_DIR / p.name).resolve()
    if not str(p).startswith(str(UPLOAD_DIR.resolve())):
        raise HTTPException(status_code=400, detail="非法路径")
    if not p.exists() or not p.is_file():
        raise HTTPException(status_code=404, detail="文件不存在")
    return StreamingResponse(
        p.open("rb"),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        # RFC 5987: filename* 的值必须 percent-encode,否则中文文件名无法进 latin-1 响应头
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(p.name)}"},
    )
